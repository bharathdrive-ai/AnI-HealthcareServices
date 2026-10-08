"""Doctors list and availability schedule: Excel workbook <-> website and portal.

    python tools/doctors_config.py check    # validate config/doctors-config.xlsx, change nothing
    python tools/doctors_config.py sync     # workbook -> portal/assets/data.js and index.html
    python tools/doctors_config.py build    # (re)create the workbook from the site's current data

The workbook is the master copy. `sync` rewrites the "DOCTORS CONFIG" blocks in data.js and
index.html; portal.js then applies the new version over each browser's saved data.
Needs: python -m pip install openpyxl
"""
import datetime as dt
import hashlib
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("openpyxl is missing. Install it with:  python -m pip install openpyxl")

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "config" / "doctors-config.xlsx"
DATA_JS = ROOT / "portal" / "assets" / "data.js"
INDEX = ROOT / "index.html"
BEGIN, END = "/* >>> DOCTORS CONFIG", "/* <<< DOCTORS CONFIG */"
NOTE = " — generated from config/doctors-config.xlsx by tools/doctors_config.py. Edit the workbook, not this block. */"

DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]
WEEK = ["Full", "AM", "PM", "Off"]
OPEN = ["Open", "Closed"]
FEE_TYPES = ["General", "Specialist"]
DESIGNATIONS = ["Professor & Head", "Professor", "Associate Professor", "Assistant Professor",
                "Senior Resident", "Junior Resident", "Consultant", "Visiting Consultant"]
EMPLOYMENT = ["Permanent", "Contract", "Visiting"]
STATUS = ["Active", "On Leave", "Inactive"]
YES_NO = ["Yes", "No"]
LEAVE_TYPES = ["Casual leave", "Sick leave", "Earned leave", "Conference leave", "Maternity leave", "Unavailable (doctor)"]
LEAVE_STATUS = ["Approved", "Pending", "Rejected"]
HEAD = "Professor & Head"
MAX_ROW = 200            # formulas and dropdowns are laid out to this row
EMAIL_DOMAIN = "ani-healthcareservices.example"

SETTINGS = [  # key, label, unit/help
    ("slotMinutes", "Slot length", "minutes per appointment slot"),
    ("amStart", "Morning session starts", "hh:mm, 24-hour"),
    ("amEnd", "Morning session ends", "hh:mm, 24-hour"),
    ("pmStart", "Afternoon session starts", "hh:mm, 24-hour"),
    ("pmEnd", "Afternoon session ends", "hh:mm, 24-hour"),
    ("maxPerSlot", "Patients per slot", "bookings allowed in one slot"),
    ("bookingWindowDays", "Booking window", "days ahead patients can book"),
    ("feeGeneral", "Consultation fee — General", "₹, for departments marked General"),
    ("feeSpecialist", "Consultation fee — Specialist", "₹, for departments marked Specialist"),
    ("feeFollowUp", "Follow-up fee", "₹, within the free follow-up period"),
    ("followUpFreeDays", "Free follow-up period", "days after a completed consultation"),
]


# ---------------------------------------------------------------- site files
def departments():
    """Departments as defined in data.js (names and codes are not edited from the workbook)."""
    rows = re.findall(r'\["(DEP\d{4})","([^"]+)","([A-Z]+)","(\w+)",(\d+),"[^"]*"\]', DATA_JS.read_text(encoding="utf-8"))
    return [{"id": i, "name": n, "code": c} for i, n, c, _, _ in rows]


def other_staff_ids():
    """IDs of non-doctor staff in data.js; doctors must not reuse them."""
    text = DATA_JS.read_text(encoding="utf-8")
    start = text.index("const staff=[")
    return set(re.findall(r'\["(STF\d{4})",', text[start:text.index("].map", start)]))


def block_span(text, path):
    a, b = text.find(BEGIN), text.find(END)
    if a < 0 or b < 0:
        sys.exit(f"Could not find the DOCTORS CONFIG markers in {path.relative_to(ROOT)}.")
    return a, b + len(END)


def read_site_config():
    text = DATA_JS.read_text(encoding="utf-8")
    a, b = block_span(text, DATA_JS)
    blk = text[a:b]
    return json.loads(blk[blk.index("{", blk.index("const DOCTORS_CONFIG=")):blk.rindex("};") + 1])


def render_data_js(cfg):
    L = [BEGIN + NOTE, "const DOCTORS_CONFIG={", f'"version":{json.dumps(cfg["version"])},', '"doctors":[']
    L += [json.dumps(d, ensure_ascii=False, separators=(",", ":")) + ("," if i < len(cfg["doctors"]) - 1 else "") for i, d in enumerate(cfg["doctors"])]
    L += ["],", f'"opdDays":{json.dumps(cfg["opdDays"], separators=(",", ":"))},', f'"generalDepts":{json.dumps(cfg["generalDepts"], separators=(",", ":"))},',
          f'"settings":{json.dumps(cfg["settings"], separators=(",", ":"))},', '"leaves":[']
    L += [json.dumps(x, ensure_ascii=False, separators=(",", ":")) + ("," if i < len(cfg["leaves"]) - 1 else "") for i, x in enumerate(cfg["leaves"])]
    L += ["]", "};", END]
    return "\n".join(L)


def render_index(cfg, depts):
    name = {d["id"]: d["name"] for d in depts}
    code = {d["id"]: d["code"] for d in depts}
    q = lambda v: json.dumps(v, ensure_ascii=False)
    L = [BEGIN + NOTE, "const DOCS=["]
    for d in cfg["doctors"]:
        if d["web"] and d["status"] != "Inactive":
            L.append(f' {{n:{q(d["name"])},q:{q(d["qual"])},s:{q(name[d["dept"]])},r:{q(d["designation"])},exp:{d["exp"]}}},')
    L += ["];", "const OPD_DAYS_BY_CODE={" + ",".join(f'{code[k]}:{json.dumps(v, separators=(",", ":"))}' for k, v in cfg["opdDays"].items()) + "};", END]
    return "\n".join(L)


def write_site(cfg):
    depts = departments()
    for path, body in ((DATA_JS, render_data_js(cfg)), (INDEX, render_index(cfg, depts))):
        text = path.read_text(encoding="utf-8")
        a, b = block_span(text, path)
        path.write_text(text[:a] + body + text[b:], encoding="utf-8", newline="\n")


def stamp(cfg):
    body = json.dumps({k: v for k, v in cfg.items() if k != "version"}, sort_keys=True, ensure_ascii=False)
    cfg["version"] = hashlib.sha1(body.encode()).hexdigest()[:12]
    return cfg


# ---------------------------------------------------------------- workbook styles
ARIAL = "Arial"
GREEN = "1F6F61"
F_HEAD = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
F_BODY = Font(name=ARIAL, size=10)
F_BOLD = Font(name=ARIAL, size=10, bold=True)
F_TITLE = Font(name=ARIAL, size=16, bold=True, color=GREEN)
F_SUB = Font(name=ARIAL, size=11, bold=True, color=GREEN)
F_NOTE = Font(name=ARIAL, size=9, italic=True, color="595959")
FILL_HEAD = PatternFill("solid", fgColor=GREEN)
FILL_INPUT = PatternFill("solid", fgColor="FFF8DC")     # cream: type here
FILL_CALC = PatternFill("solid", fgColor="EDEDED")      # grey: formula, don't edit
FILL_REF = PatternFill("solid", fgColor="E2EFDA")       # light green: reference, don't edit
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, cols, row=1):
    for i, (title, width) in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=title)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, Alignment(horizontal="center", vertical="center", wrap_text=True), BOX
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[row].height = 30


def paint(ws, rng, fill, font=F_BODY, align=None):
    for row in ws[rng]:
        for c in row:
            c.fill, c.font, c.border = fill, font, BOX
            if align:
                c.alignment = align


def dropdown(ws, rng, source, prompt=None):
    dv = DataValidation(type="list", formula1=source, allow_blank=True, showErrorMessage=True,
                        errorTitle="Pick from the list", error="Please choose a value from the dropdown list.")
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = "Tip", prompt, True
    ws.add_data_validation(dv)
    dv.add(rng)


def lst(values):
    return '"' + ",".join(values) + '"'


# ---------------------------------------------------------------- build
def build(cfg=None):
    cfg = cfg or read_site_config()
    depts = departments()
    dname = {d["id"]: d["name"] for d in depts}
    nd = len(depts)
    wb = Workbook()

    # ---- Instructions
    ws = wb.active
    ws.title = "Instructions"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 92
    ws["B2"] = "AnI-HealthcareServices — Doctors & Availability Configuration"
    ws["B2"].font = F_TITLE
    ws["B3"] = "This workbook is the master list of doctors, their weekly OPD availability, department OPD days, leave and slot settings."
    ws["B3"].font = F_NOTE
    r = 5
    ws.cell(r, 2, "Colour legend").font = F_SUB
    for fill, label, text in ((FILL_INPUT, "Cream cells", "You edit these. Most have a dropdown — pick from the list."),
                              (FILL_CALC, "Grey cells", "Calculated automatically (formulas). Don't type over them."),
                              (FILL_REF, "Green cells", "Reference only (IDs, names and codes that the website relies on). Don't change.")):
        r += 1
        c = ws.cell(r, 2, label)
        c.fill, c.font, c.border = fill, F_BOLD, BOX
        ws.cell(r, 3, text).font = F_BODY
    r += 2
    ws.cell(r, 2, "Sheets").font = F_SUB
    for name, text in (("Doctors", "One row per doctor: profile + weekly availability Mon–Sat (Full day / AM only / PM only / Off). Row order = order on the website."),
                       ("Departments", "OPD days per department (Open/Closed) and whether the consultation fee is General or Specialist. Shows head and doctor count."),
                       ("Leave", "Dates a doctor is away. Approved leave blocks that doctor's booking slots on those dates."),
                       ("Settings", "Session timings, slot length, booking window and consultation fees."),
                       ("Weekly Coverage", "Read-only: how many active doctors each department has in each morning/afternoon session. Red 0 = OPD open but nobody available.")):
        r += 1
        ws.cell(r, 2, name).font = F_BOLD
        ws.cell(r, 3, text).font = F_BODY
        ws.cell(r, 3).alignment = WRAP
    r += 2
    ws.cell(r, 2, "How to apply changes").font = F_SUB
    for i, text in enumerate(("Edit the cream cells in this workbook and save it (keep the file name and .xlsx format).",
                              "Check it:  python tools/doctors_config.py check   — lists any mistakes with the sheet and row.",
                              "Apply it:  python tools/doctors_config.py sync    — updates the website (doctor list, OPD days) and the staff portal "
                              "(staff directory, availability, leave, slot settings).",
                              "Commit and push to GitHub. Vercel and GitHub Pages redeploy automatically within a minute or two.",
                              "Each browser picks up the new configuration on its next visit. The workbook wins: changes made to these same "
                              "fields inside the portal (e.g. a doctor's weekly pattern) are replaced on the next sync."), 1):
        r += 1
        ws.cell(r, 2, f"Step {i}").font = F_BOLD
        c = ws.cell(r, 3, text)
        c.font, c.alignment = F_BODY, WRAP
        ws.row_dimensions[r].height = 28 if len(text) > 95 else 15
    r += 2
    ws.cell(r, 2, "Rules").font = F_SUB
    for text in ("Staff ID: STF + 4 digits, unique. For a new doctor use the next free ID shown below. Never reuse an old ID.",
                 "To remove a doctor, set Status = Inactive (keeps their appointment history). Inactive doctors disappear from the website and booking.",
                 "On Leave doctors stay on the website but can't be booked. For specific dates, add a row on the Leave sheet instead.",
                 "Show on website = No hides a doctor from the public 'Find a doctor' list but keeps them bookable in the portal.",
                 "Each department should have exactly one 'Professor & Head' — that name is shown as the head of department.",
                 "Emergency Care is walk-in 24×7; its OPD days don't limit emergencies.",
                 "Mobile: 10 digits starting 6–9. Email may be left blank — it is generated as firstname.lastname@" + EMAIL_DOMAIN + "."):
        r += 1
        ws.cell(r, 2, "•").font = F_BOLD
        ws.cell(r, 2).alignment = Alignment(horizontal="right")
        c = ws.cell(r, 3, text)
        c.font, c.alignment = F_BODY, WRAP
        ws.row_dimensions[r].height = 28 if len(text) > 95 else 15
    r += 2
    ws.cell(r, 2, "Next free Staff ID").font = F_BOLD
    c = ws.cell(r, 3, f'="STF"&TEXT(MAX(Doctors!$V$2:$V${MAX_ROW},20)+1,"0000")')
    c.font, c.fill, c.border = F_BOLD, FILL_CALC, BOX
    ws.cell(r + 1, 3, "IDs STF0011–STF0020 belong to nurses, technicians and other non-doctor staff.").font = F_NOTE
    r += 3
    ws.cell(r, 2, "Example — adding leave").font = F_SUB
    ex_cols = ["Staff ID", "Doctor (auto)", "Leave type", "From", "To", "Days (auto)", "Reason", "Status"]
    ex_vals = ["STF0006", "Dr. Vikram Singh", "Conference leave", dt.date(2026, 11, 16), dt.date(2026, 11, 18), 3, "Neurology conference, Delhi", "Approved"]
    r += 1
    for i, (k, v) in enumerate(zip(ex_cols, ex_vals)):  # a small vertical card (column C is wide)
        ws.cell(r + i, 2, k).font = F_BOLD
        c = ws.cell(r + i, 3, v)
        c.font = F_BODY
        c.alignment = Alignment(horizontal="left")
        if isinstance(v, dt.date):
            c.number_format = "dd-mmm-yyyy"
    ws.cell(r + len(ex_cols), 3, "This is only an example — it is not part of the configuration. Add real leave on the Leave sheet.").font = F_NOTE

    # ---- Doctors
    ws = wb.create_sheet("Doctors")
    cols = [("Staff ID", 10), ("Doctor name", 26), ("Department", 24), ("Qualification", 32), ("Designation", 21),
            ("Employment", 12), ("Experience (yrs)", 11), ("Date of joining", 13), ("Mobile", 12), ("Email", 36),
            ("Status", 10), ("Show on website", 10)] + [(d, 7) for d in DAYS] + \
           [("OPD sessions / week", 10), ("Check", 30), ("key", 4), ("id#", 4), ("dept#", 4)]
    header(ws, cols)
    for i, d in enumerate(cfg["doctors"], 2):
        vals = [d["id"], d["name"], dname[d["dept"]], d["qual"], d["designation"], d["type"], d["exp"],
                dt.date.fromisoformat(d["joined"]), d["phone"], d["email"], d["status"], "Yes" if d["web"] else "No"] + d["week"]
        for j, v in enumerate(vals, 1):
            ws.cell(i, j, v)
    last = MAX_ROW
    paint(ws, f"A2:R{last}", FILL_INPUT)
    for r in range(2, last + 1):
        ws.cell(r, 8).number_format = "dd-mmm-yyyy"
        ws.cell(r, 9).number_format = "@"
        ws.cell(r, 7).alignment = CENTER
        for c in range(11, 19):
            ws.cell(r, c).alignment = CENTER
        ws.cell(r, 19, f'=IF(A{r}="","",COUNTIF(M{r}:R{r},"Full")*2+COUNTIF(M{r}:R{r},"AM")+COUNTIF(M{r}:R{r},"PM"))')
        closed = "+".join(f'({c}{r}<>"Off")*(INDEX(Departments!${d}$2:${d}${nd + 1},$W{r})="Closed")'
                          for c, d in zip("MNOPQR", "DEFGHI"))
        ws.cell(r, 20, f'=IF(A{r}="","",IF(COUNTIF($A$2:$A${last},A{r})>1,"Duplicate Staff ID",'
                       f'IF(COUNTA(B{r}:I{r})+COUNTA(K{r}:L{r})<10,"Fill in all details",'
                       f'IF($W{r}="","Pick a department",'
                       f'IF(COUNTIF(M{r}:R{r},"Off")=6,"No OPD days set",'
                       f'IF({closed}>0,"Note: works on a closed OPD day","OK"))))))')
        ws.cell(r, 21, f'=C{r}&"|"&E{r}')
        ws.cell(r, 22, f'=IF(A{r}="","",IFERROR(VALUE(MID(A{r},4,4)),0))')
        ws.cell(r, 23, f'=IFERROR(MATCH(C{r},Departments!$B$2:$B${nd + 1},0),"")')
    paint(ws, f"S2:T{last}", FILL_CALC)
    for r in range(2, last + 1):
        ws.cell(r, 19).alignment = CENTER
    paint(ws, f"U2:W{last}", FILL_CALC, F_NOTE)
    for col in "UVW":
        ws.column_dimensions[col].hidden = True  # helper columns for the formulas
    dropdown(ws, f"C2:C{last}", f"=Departments!$B$2:$B${nd + 1}")
    dropdown(ws, f"E2:E{last}", lst(DESIGNATIONS))
    dropdown(ws, f"F2:F{last}", lst(EMPLOYMENT))
    dropdown(ws, f"K2:K{last}", lst(STATUS), "Inactive = removed from the website and booking (history is kept).")
    dropdown(ws, f"L2:L{last}", lst(YES_NO))
    dropdown(ws, f"M2:R{last}", lst(WEEK), "Full = 09:00–13:00 and 14:00–16:00 · AM = morning only · PM = afternoon only · Off")
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="60", allow_blank=True, showErrorMessage=True,
                        error="Experience must be a whole number of years (0–60).")
    ws.add_data_validation(dv)
    dv.add(f"G2:G{last}")
    dv = DataValidation(type="date", operator="between", formula1="DATE(1960,1,1)", formula2="DATE(2100,12,31)", allow_blank=True,
                        showErrorMessage=True, error="Enter a date, e.g. 01-Jul-2024.")
    ws.add_data_validation(dv)
    dv.add(f"H2:H{last}")
    rng = f"M2:R{last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Full"'], fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=ARIAL, color="006100")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"AM"'], fill=PatternFill("solid", fgColor="DDEBF7"), font=Font(name=ARIAL, color="1F4E79")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PM"'], fill=PatternFill("solid", fgColor="FCE4D6"), font=Font(name=ARIAL, color="843C0C")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Off"'], fill=PatternFill("solid", fgColor="D9D9D9"), font=Font(name=ARIAL, color="595959")))
    ws.conditional_formatting.add(f"T2:T{last}", FormulaRule(formula=['AND(T2<>"",T2<>"OK",LEFT(T2,4)<>"Note")'], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=ARIAL, bold=True, color="9C0006")))
    ws.conditional_formatting.add(f"K2:K{last}", CellIsRule(operator="equal", formula=['"Inactive"'], font=Font(name=ARIAL, color="9C0006")))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:T{last}"
    ws["M1"].comment = Comment("Weekly availability Mon–Sat. Full = whole day, AM = morning session, PM = afternoon session, Off = no OPD.", "Config")
    ws["L1"].comment = Comment("No = hidden from the public 'Find a doctor' list (still bookable from the portal).", "Config")

    # ---- Departments
    ws = wb.create_sheet("Departments")
    header(ws, [("Dept ID", 10), ("Department", 26), ("Code", 7)] + [(d, 8) for d in DAYS] +
               [("Consultation fee", 13), ("Active doctors", 10), ("Head of department", 26)])
    for i, d in enumerate(depts, 2):
        ws.cell(i, 1, d["id"])
        ws.cell(i, 2, d["name"])
        ws.cell(i, 3, d["code"])
        for j, v in enumerate(cfg["opdDays"].get(d["id"], [1] * 6)):
            ws.cell(i, 4 + j, "Open" if v else "Closed")
        ws.cell(i, 10, "General" if d["id"] in cfg["generalDepts"] else "Specialist")
        ws.cell(i, 11, f'=COUNTIFS(Doctors!$C$2:$C${MAX_ROW},B{i},Doctors!$K$2:$K${MAX_ROW},"Active")')
        ws.cell(i, 12, f'=IFERROR(INDEX(Doctors!$B$2:$B${MAX_ROW},MATCH(B{i}&"|{HEAD}",Doctors!$U$2:$U${MAX_ROW},0)),"— none —")')
    e = nd + 1
    paint(ws, f"A2:C{e}", FILL_REF)
    paint(ws, f"D2:J{e}", FILL_INPUT, align=CENTER)
    paint(ws, f"K2:L{e}", FILL_CALC)
    for r in range(2, e + 1):
        ws.cell(r, 11).alignment = CENTER
    dropdown(ws, f"D2:I{e}", lst(OPEN))
    dropdown(ws, f"J2:J{e}", lst(FEE_TYPES))
    ws.conditional_formatting.add(f"D2:I{e}", CellIsRule(operator="equal", formula=['"Closed"'], fill=PatternFill("solid", fgColor="D9D9D9"), font=Font(name=ARIAL, color="595959")))
    ws.conditional_formatting.add(f"L2:L{e}", CellIsRule(operator="equal", formula=['"— none —"'], font=Font(name=ARIAL, bold=True, color="9C0006")))
    ws.freeze_panes = "C2"
    ws.cell(e + 2, 2, "OPD runs Monday–Saturday; Sunday is always closed. Department names and codes are fixed (the website uses them).").font = F_NOTE
    ws.cell(e + 3, 2, "Emergency Care is walk-in 24×7 — its OPD days do not limit emergency care.").font = F_NOTE

    # ---- Leave
    ws = wb.create_sheet("Leave")
    header(ws, [("Staff ID", 10), ("Doctor (auto)", 26), ("Leave type", 20), ("From", 13), ("To", 13), ("Days (auto)", 9), ("Reason", 36), ("Status", 11)])
    for i, x in enumerate(cfg["leaves"], 2):
        for j, v in enumerate([x["staff"], None, x["type"], dt.date.fromisoformat(x["from"]), dt.date.fromisoformat(x["to"]), None, x["reason"], x["status"]], 1):
            if v is not None:
                ws.cell(i, j, v)
    for r in range(2, MAX_ROW + 1):
        ws.cell(r, 2, f'=IF(A{r}="","",IFERROR(INDEX(Doctors!$B$2:$B${MAX_ROW},MATCH(A{r},Doctors!$A$2:$A${MAX_ROW},0)),"Unknown Staff ID"))')
        ws.cell(r, 6, f'=IF(OR(D{r}="",E{r}=""),"",E{r}-D{r}+1)')
    paint(ws, f"A2:A{MAX_ROW}", FILL_INPUT)
    paint(ws, f"C2:E{MAX_ROW}", FILL_INPUT)
    paint(ws, f"G2:H{MAX_ROW}", FILL_INPUT)
    paint(ws, f"B2:B{MAX_ROW}", FILL_CALC)
    paint(ws, f"F2:F{MAX_ROW}", FILL_CALC, align=CENTER)
    for r in range(2, MAX_ROW + 1):
        ws.cell(r, 4).number_format = ws.cell(r, 5).number_format = "dd-mmm-yyyy"
        ws.cell(r, 6).number_format = "0"
    dropdown(ws, f"A2:A{MAX_ROW}", f"=Doctors!$A$2:$A${MAX_ROW}", "Pick the doctor's Staff ID; the name fills in automatically.")
    dropdown(ws, f"C2:C{MAX_ROW}", lst(LEAVE_TYPES))
    dropdown(ws, f"H2:H{MAX_ROW}", lst(LEAVE_STATUS), "Only Approved leave blocks booking slots.")
    dv = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2100,12,31)", allow_blank=True,
                        showErrorMessage=True, error="Enter a date, e.g. 16-Nov-2026.")
    ws.add_data_validation(dv)
    dv.add(f"D2:E{MAX_ROW}")
    ws.conditional_formatting.add(f"B2:B{MAX_ROW}", CellIsRule(operator="equal", formula=['"Unknown Staff ID"'], font=Font(name=ARIAL, bold=True, color="9C0006")))
    ws.conditional_formatting.add(f"F2:F{MAX_ROW}", CellIsRule(operator="lessThan", formula=["1"], fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.freeze_panes = "A2"

    # ---- Settings
    ws = wb.create_sheet("Settings")
    header(ws, [("Setting", 32), ("Value", 12), ("Notes", 44)])
    s = cfg["settings"]
    vals = {"slotMinutes": s["slotMinutes"], "amStart": s["AM"]["start"], "amEnd": s["AM"]["end"], "pmStart": s["PM"]["start"],
            "pmEnd": s["PM"]["end"], "maxPerSlot": s["maxPerSlot"], "bookingWindowDays": s["bookingWindowDays"],
            "feeGeneral": s["fees"]["general"], "feeSpecialist": s["fees"]["specialist"], "feeFollowUp": s["fees"]["followUp"],
            "followUpFreeDays": s["followUpFreeDays"]}
    for i, (k, label, note) in enumerate(SETTINGS, 2):
        ws.cell(i, 1, label)
        c = ws.cell(i, 2, vals[k])
        if ":" in str(vals[k]):
            c.number_format = "@"
        ws.cell(i, 3, note)
    e = len(SETTINGS) + 1
    paint(ws, f"A2:A{e}", FILL_REF, F_BOLD)
    paint(ws, f"B2:B{e}", FILL_INPUT, align=CENTER)
    paint(ws, f"C2:C{e}", PatternFill(fill_type=None), F_NOTE)
    ws.cell(e + 2, 1, "Times are text in 24-hour hh:mm form (e.g. 09:00, 14:30). Slot length should divide each session evenly.").font = F_NOTE

    # ---- Weekly Coverage
    ws = wb.create_sheet("Weekly Coverage")
    ws["A1"] = "Active doctors available per session (calculated — no need to edit)"
    ws["A1"].font = F_SUB
    ws.merge_cells("B2:G2")
    ws.merge_cells("H2:M2")
    ws["B2"], ws["H2"] = "Morning (AM)", "Afternoon (PM)"
    for c in ("B2", "H2"):
        ws[c].font, ws[c].fill, ws[c].alignment = F_HEAD, FILL_HEAD, CENTER
    header(ws, [("Department", 26)] + [(d, 8) for d in DAYS] * 2, row=3)
    for i in range(nd):
        r, src = 4 + i, 2 + i
        ws.cell(r, 1, f"=Departments!B{src}")
        for j in range(6):
            dcol = "DEFGHI"[j]
            mcol = "MNOPQR"[j]
            for k, sess in enumerate(("AM", "PM")):
                f = (f'=IF(Departments!{dcol}{src}="Closed","Closed",'
                     f'COUNTIFS(Doctors!$C$2:$C${MAX_ROW},$A{r},Doctors!${mcol}$2:${mcol}${MAX_ROW},"Full",Doctors!$K$2:$K${MAX_ROW},"Active")+'
                     f'COUNTIFS(Doctors!$C$2:$C${MAX_ROW},$A{r},Doctors!${mcol}$2:${mcol}${MAX_ROW},"{sess}",Doctors!$K$2:$K${MAX_ROW},"Active"))')
                ws.cell(r, 2 + j + 6 * k, f)
    e = 3 + nd
    paint(ws, f"A4:A{e}", FILL_CALC, F_BOLD)
    paint(ws, f"B4:M{e}", FILL_CALC, align=CENTER)
    ws.conditional_formatting.add(f"B4:M{e}", CellIsRule(operator="equal", formula=["0"], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=ARIAL, bold=True, color="9C0006")))
    ws.conditional_formatting.add(f"B4:M{e}", CellIsRule(operator="equal", formula=['"Closed"'], font=Font(name=ARIAL, color="808080")))
    ws.cell(e + 2, 1, "Counts doctors with Status = Active whose weekly pattern covers that session. Leave dates are not included here.").font = F_NOTE
    ws.freeze_panes = "B4"

    for sheet in wb.worksheets:
        sheet.sheet_properties.tabColor = GREEN if sheet.title in ("Doctors", "Departments", "Leave", "Settings") else "A6A6A6"
    wb.calculation.fullCalcOnLoad = True
    XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX)
    print(f"Wrote {XLSX.relative_to(ROOT)} — {len(cfg['doctors'])} doctors, {nd} departments, {len(cfg['leaves'])} leave rows.")


# ---------------------------------------------------------------- read + validate
def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    t = _text(v)
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d %b %Y"):
        try:
            return dt.datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    return None


def _time(v):
    if isinstance(v, (dt.time, dt.datetime)):
        return v.strftime("%H:%M")
    if isinstance(v, (int, float)) and 0 <= v < 1:  # Excel day fraction
        m = round(v * 1440)
        return f"{m // 60:02d}:{m % 60:02d}"
    m = re.fullmatch(r"(\d{1,2})[:.](\d{2})", _text(v))
    if m and int(m[1]) < 24 and int(m[2]) < 60:
        return f"{int(m[1]):02d}:{m[2]}"
    return None


def _int(v):
    t = _text(v)
    return int(t) if re.fullmatch(r"\d+", t) else None


def read_workbook():
    if not XLSX.exists():
        sys.exit(f"{XLSX.relative_to(ROOT)} not found. Create it with:  python tools/doctors_config.py build")
    wb = load_workbook(XLSX)
    for name in ("Doctors", "Departments", "Leave", "Settings"):
        if name not in wb.sheetnames:
            sys.exit(f"Sheet '{name}' is missing from the workbook.")
    depts = departments()
    by_name = {d["name"].lower(): d for d in depts}
    errors, warnings = [], []
    today = dt.date.today()
    reserved = other_staff_ids()

    # Departments
    opd, general = {}, []
    ws = wb["Departments"]
    for r in range(2, ws.max_row + 1):
        did = _text(ws.cell(r, 1).value)
        if not did.startswith("DEP"):
            continue
        if did not in {d["id"] for d in depts}:
            errors.append(f"Departments row {r}: unknown Dept ID {did} (department IDs can't be added here).")
            continue
        days = []
        for j, day in enumerate(DAYS):
            v = _text(ws.cell(r, 4 + j).value).capitalize()
            if v not in OPEN:
                errors.append(f"Departments row {r} ({did}) {day}: '{v}' must be Open or Closed.")
            days.append(1 if v == "Open" else 0)
        opd[did] = days
        fee = _text(ws.cell(r, 10).value).capitalize()
        if fee not in FEE_TYPES:
            errors.append(f"Departments row {r} ({did}): Consultation fee must be General or Specialist.")
        elif fee == "General":
            general.append(did)
    for d in depts:
        if d["id"] not in opd:
            errors.append(f"Departments: {d['id']} {d['name']} is missing — don't delete department rows.")

    # Doctors
    doctors, seen = [], {}
    ws = wb["Doctors"]
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, 19)]
        if not any(_text(v) for v in row):
            continue
        where = f"Doctors row {r}"
        did, name = _text(row[0]).upper(), _text(row[1])
        if not re.fullmatch(r"STF\d{4}", did):
            errors.append(f"{where}: Staff ID '{row[0] or ''}' must look like STF0047.")
        elif did in seen:
            errors.append(f"{where}: Staff ID {did} is already used on row {seen[did]}.")
        elif did in reserved:
            errors.append(f"{where}: Staff ID {did} belongs to a non-doctor staff member — use the next free ID.")
        seen.setdefault(did, r)
        if not name:
            errors.append(f"{where}: Doctor name is empty.")
        elif not name.startswith("Dr."):
            warnings.append(f"{where}: '{name}' doesn't start with 'Dr.'.")
        dep = by_name.get(_text(row[2]).lower())
        if not dep:
            errors.append(f"{where}: Department '{_text(row[2])}' is not in the Departments list.")
        qual, desig = _text(row[3]), _text(row[4])
        if not qual:
            errors.append(f"{where}: Qualification is empty.")
        if not desig:
            errors.append(f"{where}: Designation is empty.")
        elif desig not in DESIGNATIONS:
            warnings.append(f"{where}: Designation '{desig}' is not one of the standard titles.")
        emp = _text(row[5]).capitalize()
        if emp not in EMPLOYMENT:
            errors.append(f"{where}: Employment must be one of {', '.join(EMPLOYMENT)}.")
        exp = _int(row[6])
        if exp is None or exp > 60:
            errors.append(f"{where}: Experience must be a whole number of years.")
        joined = _date(row[7])
        if not joined:
            errors.append(f"{where}: Date of joining is missing or not a date.")
        elif joined > today:
            errors.append(f"{where}: Date of joining {joined} is in the future.")
        phone = re.sub(r"[\s-]", "", _text(row[8]))
        if not re.fullmatch(r"[6-9]\d{9}", phone):
            errors.append(f"{where}: Mobile '{_text(row[8])}' must be 10 digits starting 6–9.")
        email = _text(row[9]).lower()
        if not email and name:
            email = re.sub(r"\s+", ".", re.sub(r"^(dr|sr)\.\s*", "", name.lower())) + "@" + EMAIL_DOMAIN
        if email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            errors.append(f"{where}: Email '{email}' is not valid.")
        status = {s.lower(): s for s in STATUS}.get(_text(row[10]).lower())
        if not status:
            errors.append(f"{where}: Status must be one of {', '.join(STATUS)}.")
        web = _text(row[11]).capitalize()
        if web not in YES_NO:
            errors.append(f"{where}: Show on website must be Yes or No.")
        week = []
        for j, day in enumerate(DAYS):
            v = {w.lower(): w for w in WEEK}.get(_text(row[12 + j]).lower())
            if not v:
                errors.append(f"{where} {day}: '{_text(row[12 + j])}' must be Full, AM, PM or Off.")
            week.append(v or "Off")
        if dep and status == "Active":
            closed = [DAYS[j] for j in range(6) if week[j] != "Off" and not opd.get(dep["id"], [1] * 6)[j]]
            if closed and dep["code"] != "EMR":
                warnings.append(f"{where} {name}: scheduled on {', '.join(closed)} but {dep['name']} OPD is closed then (no slots will show).")
        doctors.append({"id": did, "name": name, "dept": dep["id"] if dep else "", "qual": qual, "designation": desig, "type": emp,
                        "exp": exp or 0, "joined": joined.isoformat() if joined else "", "phone": phone, "email": email,
                        "status": status or "Active", "web": web == "Yes", "week": week})
    if not doctors:
        errors.append("Doctors: the sheet has no doctors.")
    for d in depts:
        heads = [x["name"] for x in doctors if x["dept"] == d["id"] and x["designation"] == HEAD and x["status"] != "Inactive"]
        if len(heads) != 1:
            warnings.append(f"{d['name']}: {'no' if not heads else len(heads)} 'Professor & Head' ({', '.join(heads) or '—'}).")
        if d["code"] != "EMR":
            for j, day in enumerate(DAYS):
                if opd.get(d["id"], [0] * 6)[j]:
                    for sess in ("AM", "PM"):
                        if not any(x["dept"] == d["id"] and x["status"] == "Active" and x["week"][j] in ("Full", sess) for x in doctors):
                            warnings.append(f"{d['name']}: OPD open on {day} but no active doctor in the {sess} session.")

    # Leave
    leaves, ids = [], {d["id"]: d for d in doctors}
    ws = wb["Leave"]
    for r in range(2, ws.max_row + 1):
        sid, ltype, f, t, reason, st = (ws.cell(r, c).value for c in (1, 3, 4, 5, 7, 8))
        if not any(_text(v) for v in (sid, ltype, f, t, reason, st)):
            continue
        where = f"Leave row {r}"
        sid = _text(sid).upper()
        if sid not in ids:
            errors.append(f"{where}: Staff ID '{sid}' is not on the Doctors sheet.")
        ltype = _text(ltype)
        if ltype not in LEAVE_TYPES:
            errors.append(f"{where}: Leave type must be one of {', '.join(LEAVE_TYPES)}.")
        fd, td = _date(f), _date(t)
        if not fd or not td:
            errors.append(f"{where}: From and To must be dates.")
        elif td < fd:
            errors.append(f"{where}: To ({td}) is before From ({fd}).")
        st = _text(st).capitalize()
        if st not in LEAVE_STATUS:
            errors.append(f"{where}: Status must be one of {', '.join(LEAVE_STATUS)}.")
        if fd and td and td >= fd:
            lid = f"LVC-{sid}-{fd.isoformat()}"
            if any(x["id"] == lid for x in leaves):
                errors.append(f"{where}: {sid} already has leave starting {fd}.")
            leaves.append({"id": lid, "staff": sid, "type": ltype, "from": fd.isoformat(), "to": td.isoformat(),
                           "days": (td - fd).days + 1, "reason": _text(reason) or ltype, "status": st})

    # Settings
    ws = wb["Settings"]
    raw = {}
    labels = {label: k for k, label, _ in SETTINGS}
    for r in range(2, ws.max_row + 1):
        k = labels.get(_text(ws.cell(r, 1).value))
        if k:
            raw[k] = ws.cell(r, 2).value
    st = {}
    for k, label, _ in SETTINGS:
        if k not in raw:
            errors.append(f"Settings: '{label}' row is missing.")
            continue
        v = _time(raw[k]) if k.endswith(("Start", "End")) else _int(raw[k])
        if v is None:
            errors.append(f"Settings: '{label}' value '{_text(raw[k])}' is not valid ({'hh:mm' if k.endswith(('Start', 'End')) else 'whole number'}).")
        st[k] = v
    if not errors:
        mins = lambda h: int(h[:2]) * 60 + int(h[3:])
        if st["slotMinutes"] < 5:
            errors.append("Settings: Slot length must be at least 5 minutes.")
        if st["maxPerSlot"] < 1 or st["bookingWindowDays"] < 1:
            errors.append("Settings: Patients per slot and Booking window must be at least 1.")
        if not mins(st["amStart"]) < mins(st["amEnd"]) <= mins(st["pmStart"]) < mins(st["pmEnd"]):
            errors.append("Settings: sessions must run morning start < morning end ≤ afternoon start < afternoon end.")
        for a, b in (("amStart", "amEnd"), ("pmStart", "pmEnd")):
            if not errors and (mins(st[b]) - mins(st[a])) % st["slotMinutes"]:
                warnings.append(f"Settings: {st[a]}–{st[b]} isn't a whole number of {st['slotMinutes']}-minute slots; the last part is unused.")
    settings = {} if errors else {
        "slotMinutes": st["slotMinutes"], "AM": {"start": st["amStart"], "end": st["amEnd"]}, "PM": {"start": st["pmStart"], "end": st["pmEnd"]},
        "maxPerSlot": st["maxPerSlot"], "bookingWindowDays": st["bookingWindowDays"],
        "fees": {"general": st["feeGeneral"], "specialist": st["feeSpecialist"], "followUp": st["feeFollowUp"]},
        "followUpFreeDays": st["followUpFreeDays"]}
    cfg = {"version": "", "doctors": doctors, "opdDays": {d["id"]: opd.get(d["id"], [1] * 6) for d in depts},
           "generalDepts": general, "settings": settings, "leaves": leaves}
    return cfg, errors, warnings


def report(cfg, errors, warnings):
    for w in warnings:
        print("  warning:", w)
    for e in errors:
        print("  ERROR:  ", e)
    active = sum(d["status"] == "Active" for d in cfg["doctors"])
    web = sum(d["web"] and d["status"] != "Inactive" for d in cfg["doctors"])
    print(f"{len(cfg['doctors'])} doctors ({active} active, {web} on website), {len(cfg['leaves'])} leave rows, "
          f"{len(errors)} error(s), {len(warnings)} warning(s).")


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "build":
        if XLSX.exists() and "--force" not in sys.argv:
            sys.exit(f"{XLSX.relative_to(ROOT)} already exists. Add --force to overwrite it with the site's current data.")
        build()
    elif cmd in ("check", "sync"):
        cfg, errors, warnings = read_workbook()
        report(cfg, errors, warnings)
        if errors:
            sys.exit("Fix the errors above in the workbook, save it, and run again. Nothing was changed.")
        if cmd == "sync":
            stamp(cfg)
            if cfg["version"] == read_site_config().get("version"):
                print("The website already matches the workbook — nothing to update.")
                return
            write_site(cfg)
            print(f"Updated portal/assets/data.js and index.html (config version {cfg['version']}). Commit and push to publish.")
    else:
        print(__doc__)
        sys.exit(1)


if __name__ == "__main__":
    main()
